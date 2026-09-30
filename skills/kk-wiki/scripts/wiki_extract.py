# -*- coding: utf-8 -*-
"""kk-wiki 첨부·문서 텍스트 추출 — 위키 첨부(attachments/<pageId>/<파일>)와 외부 문서(docs/<파일>)를 절 단위 .md 로.

  python wiki_extract.py run [--force] [--only 문자열] [--ext hwp,zip]   # → attachments_text/<pageId|docs>/<파일>/NN_<절>.md + attachments_index.json/.md
  python wiki_extract.py status                          # 파일 수·절 수·추출 실패 목록
  python wiki_extract.py show <파일 문자열>               # 그 파일의 절 목록(제목·쪽·글자수·tier·대체 페이지)
  python wiki_extract.py restale                         # 규정 변경표를 고친 뒤 '⚠ 구값' 표시만 다시(다시 추출하지 않음, 몇 초)

형식: pdf(PyMuPDF) · hwp(hwp5html, pyhwp — 표 칸까지; 실패하면 hwp5txt) · hwpx/docx(문단·표를 문서 순서대로)/pptx/xlsx/xls(zip·python-docx·openpyxl·python-pptx·xlrd) ·
doc(Word COM, 있을 때만) · zip(안의 파일을 안쪽 경로 그대로 식별해 같은 규칙으로 — a/같은이름, b/같은이름 도 따로) ·
그림은 목록만. 폰트 인코딩 때문에 글자가 깨진 PDF 는 `garbled` 로 표시하고 같은 이름의 hwpx/docx 가 있으면 그쪽을 쓴다.

절마다 frontmatter: tier(1 현행 / 2 첨부 / 3 구버전) · section_kind(rule/procedure/case/table/reference) · doc_date · pages · superseded_by(같은 내용을 옮겨 적은
위키 페이지 — 제목·본문 겹침으로 자동 판정 + docs_manifest 의 규칙) · stale_values(`skills/_shared/rule_changes.md` 의 옛 값이 들어 있으면).
선별 규칙(제외 쪽·문서 날짜·대체 페이지)은 `assets/docs_manifest.json`(기본) + `<스냅샷>/docs_manifest.json`(덮어씀).
개인정보(메일 주소·휴대전화·주민번호 꼴)는 추출 때 지운다. 결과는 KIST 내부 자료 — `{kiki_root}/wiki/` 밖으로 내보내지 않는다.
"""
from __future__ import annotations
import argparse, datetime, hashlib, io, json, os, re, shutil, subprocess, sys, tempfile, zipfile
from html import unescape

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from wiki_snapshot import snapshot_root, safe_seg  # noqa: E402

ASSET_MANIFEST = os.path.join(HERE, "..", "assets", "docs_manifest.json")
RULE_CHANGES = os.path.join(HERE, "..", "..", "_shared", "rule_changes.md")
TEXT_DIR = "attachments_text"
INDEX_JSON = "attachments_index.json"
INDEX_MD = "attachments_index.md"
MAX_SECTION = 9000      # 절 하나의 글자 상한(넘으면 문단 경계에서 자름)
MIN_SECTION = 1500      # 이보다 짧으면 다음 머리글에서 자르지 않고 이어 붙임
OLD_YEARS = 3           # doc_date 가 이보다 오래되면 tier 3
TODAY = datetime.date.today()

SUPPORTED = {"pdf", "hwp", "hwpx", "docx", "pptx", "xlsx", "xls", "doc", "zip", "txt", "md", "csv"}
LIST_ONLY = {"png", "jpg", "jpeg", "gif", "bmp"}


# ---------------------------------------------------------------- 공통
def log(s: str) -> None:
    print(s, flush=True)


def load_manifest(root: str) -> dict:
    m = {"docs": []}
    for p in (ASSET_MANIFEST, os.path.join(root, "docs_manifest.json")):
        if os.path.exists(p):
            try:
                d = json.load(io.open(p, encoding="utf-8-sig"))
            except Exception as e:
                log(f"[kk-wiki] manifest 읽기 실패 {p}: {e}")
                continue
            for doc in d.get("docs", []):
                m["docs"] = [x for x in m["docs"] if x.get("match") != doc.get("match")] + [doc]
            for k, v in d.items():
                if k != "docs":
                    m[k] = v
    return m


def manifest_for(m: dict, fname: str) -> dict:
    for doc in m.get("docs", []):
        if doc.get("match") and doc["match"] in fname:
            return doc
    return {}


def scrub(t: str) -> str:
    """개인정보 꼴 제거 — 메일 주소·휴대전화·주민번호. 내선 4자리는 위키 본문과 같은 수준이라 둔다."""
    t = re.sub(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", "[메일]", t)
    t = re.sub(r"01[016789][-.\s]?\d{3,4}[-.\s]?\d{4}", "[휴대전화]", t)
    t = re.sub(r"(?<!\d)\d{6}[-\s]?[1-4]\d{6}(?!\d)", "[주민번호]", t)
    return t


def garbled_ratio(t: str) -> float:
    """폰트 인코딩이 깨진 PDF 판정 — 한글 음절·영숫자·흔한 기호가 아닌 글자 비율."""
    s = re.sub(r"\s", "", t)
    if len(s) < 200:
        return 0.0
    bad = sum(1 for ch in s if not (
        "가" <= ch <= "힣" or ch.isascii() or "　" <= ch <= "〿" or "＀" <= ch <= "￯"
        or " " <= ch <= "⯿" or "一" <= ch <= "鿿" or ch in "·※○●□■◆◇▶▷★☆☎℃㎡㎏㎜㎝㎞㎖㎥"))
    return bad / len(s)


# ---------------------------------------------------------------- 형식별 추출 → pages(list[str])
def ext_pdf(path: str) -> tuple[list, dict]:
    import fitz
    d = fitz.open(path)
    pages = [pg.get_text("text") for pg in d]
    meta = {"pdf_pages": len(d), "pdf_create": (d.metadata or {}).get("creationDate", ""), "toc": d.get_toc() or []}
    return pages, meta


def _hwp_txt(path: str) -> str:
    """hwp5txt — 빠르지만 표 안 글자를 `<표>` 로 뭉갠다(양식·전결표가 비어 나옴). html 이 실패했을 때만."""
    exe = shutil.which("hwp5txt") or shutil.which("hwp5txt.exe")
    if not exe:
        raise RuntimeError("hwp5txt 없음 — `pip install pyhwp`")
    fd, tmp = tempfile.mkstemp(suffix=".txt"); os.close(fd)
    try:
        r = subprocess.run([exe, "--output", tmp, path], capture_output=True, timeout=180)
        if r.returncode != 0:
            err = (r.stderr or b"").decode("utf-8", "replace").strip().splitlines()
            raise RuntimeError("hwp5txt 실패: " + (err[-1][:100] if err else f"code {r.returncode}"))
        return io.open(tmp, encoding="utf-8", errors="replace").read()
    finally:
        try:
            os.remove(tmp)
        except Exception:
            pass


def ext_hwp(path: str) -> tuple[list, dict]:
    """hwp → hwp5html(pyhwp) 의 xhtml 에서 글자를 뽑는다 — 표의 칸도 `|` 로 이어져 나온다(hwp5txt 는 표를 `<표>` 로 뭉갬)."""
    exe = shutil.which("hwp5html") or shutil.which("hwp5html.exe")
    if not exe:
        return [_hwp_txt(path)], {"hwp_via": "txt"}
    tmpd = tempfile.mkdtemp(prefix="kkwiki_hwp_")
    try:
        try:
            r = subprocess.run([exe, "--output", tmpd, path], capture_output=True, timeout=300)
        except subprocess.TimeoutExpired:
            return [_hwp_txt(path)], {"hwp_via": "txt(html 시간 초과)"}
        xp = os.path.join(tmpd, "index.xhtml")
        if r.returncode != 0 or not os.path.exists(xp):
            try:
                return [_hwp_txt(path)], {"hwp_via": "txt"}
            except Exception as e2:
                err = (r.stderr or b"").decode("utf-8", "replace").strip().splitlines()
                raise RuntimeError("hwp5html 실패: " + (err[-1][:80] if err else f"code {r.returncode}") + f" / {str(e2)[:60]}")
        x = io.open(xp, encoding="utf-8", errors="replace").read()
    finally:
        shutil.rmtree(tmpd, ignore_errors=True)
    x = re.sub(r"<(style|script|head).*?</\1>", "", x, flags=re.S | re.I)
    x = re.sub(r"<br\s*/?>", "\n", x, flags=re.I)
    x = re.sub(r"</(p|div|h\d|tr|li|table)>", "\n", x, flags=re.I)
    x = re.sub(r"</t[dh]>", " | ", x, flags=re.I)
    t = unescape(_TAG.sub("", x)).replace("\r", "")
    t = re.sub(r"[ \t ]+", " ", t)
    t = re.sub(r"\n\s*\| ", " | ", t)               # 칸 안 단락 줄바꿈 → 한 행에 `칸 | 칸`
    t = re.sub(r"( \| )+\n", "\n", t)              # 빈 칸만 남은 행 꼬리
    t = re.sub(r"\n\s*\n+", "\n\n", t).strip()
    return [t], {"hwp_via": "html"}


_TAG = re.compile(r"<[^>]+>")


def _xml_text(xml: str, para_end: str, cell_end: str = "") -> str:
    xml = re.sub(r"<[^>]*?/>", "", xml)  # 자기닫힘 태그 제거(줄바꿈 태그는 아래에서)
    xml = xml.replace(para_end, "\n")
    if cell_end:
        xml = xml.replace(cell_end, " | ")
    return unescape(_TAG.sub("", xml))


def ext_hwpx(path: str) -> tuple[list, dict]:
    out = []
    with zipfile.ZipFile(path) as z:
        names = sorted([n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n)],
                       key=lambda n: int(re.search(r"(\d+)", n).group(1)))
        for n in names:
            xml = z.read(n).decode("utf-8", "replace")
            xml = re.sub(r"<hp:lineBreak\s*/>", "\n", xml)
            out.append(_xml_text(xml, "</hp:p>", "</hp:tc>"))
    return ["\n".join(out)], {}


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def ext_docx(path: str) -> tuple[list, dict]:
    """본문을 **문서 순서대로** — 문단과 표가 섞인 그대로(Codex 검토 v0.7.8: 문단을 다 뽑고 표를 끝에 몰아 붙여 순서가 틀렸다).
    가로로 병합된 칸은 한 번만, 내용 컨트롤(sdt, 목차 등) 안 문단도 읽는다."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    lines = []

    def table_lines(t) -> list:
        out = []
        for row in t.rows:
            cells, prev = [], None
            for c in row.cells:
                if prev is not None and c._tc is prev:
                    continue
                prev = c._tc
                cells.append(c.text.strip().replace("\n", " "))
            out.append("| " + " | ".join(cells) + " |")
        return out

    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(Paragraph(el, d).text)
        elif tag == "tbl":
            lines += [""] + table_lines(Table(el, d)) + [""]
        elif tag == "sdt":
            for p in el.iter(W_NS + "p"):
                lines.append("".join(t.text or "" for t in p.iter(W_NS + "t")))
    return ["\n".join(lines)], {}


def ext_pptx(path: str) -> tuple[list, dict]:
    from pptx import Presentation
    pr = Presentation(path)
    pages = []
    for i, s in enumerate(pr.slides, 1):
        buf = [f"## 슬라이드 {i}"]
        for sh in s.shapes:
            if sh.has_text_frame:
                buf.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    buf.append(" | ".join(c.text.strip() for c in row.cells))
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            nt = s.notes_slide.notes_text_frame.text.strip()
            if nt:
                buf.append("(메모) " + nt)
        pages.append("\n".join(buf))
    return pages, {}


def ext_xlsx(path: str) -> tuple[list, dict]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pages = []
    for ws in wb.worksheets:
        buf = [f"## 시트: {ws.title}"]
        n = 0
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip().replace("\n", " ") for v in row]
            if any(cells):
                buf.append(" | ".join(cells).rstrip(" |"))
                n += 1
                if n >= 4000:
                    buf.append("(… 4000행에서 끊음)"); break
        pages.append("\n".join(buf))
    return pages, {}


def ext_xls(path: str) -> tuple[list, dict]:
    import xlrd
    wb = xlrd.open_workbook(path)
    pages = []
    for ws in wb.sheets():
        buf = [f"## 시트: {ws.name}"]
        for r in range(min(ws.nrows, 4000)):
            cells = [str(ws.cell_value(r, c)).strip().replace("\n", " ") for c in range(ws.ncols)]
            if any(cells):
                buf.append(" | ".join(cells).rstrip(" |"))
        pages.append("\n".join(buf))
    return pages, {}


def ext_doc(path: str) -> tuple[list, dict]:
    try:
        import win32com.client as w
    except Exception:
        raise RuntimeError("doc 은 Word(COM) 가 있을 때만")
    word = w.Dispatch("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(os.path.abspath(path), ReadOnly=True, ConfirmConversions=False, AddToRecentFiles=False)
        try:
            t = d.Content.Text
        finally:
            d.Close(False)
    finally:
        word.Quit()
    return [t.replace("\r\x07", " | ").replace("\r", "\n")], {}


def ext_text(path: str) -> tuple[list, dict]:
    return [io.open(path, encoding="utf-8-sig", errors="replace").read()], {}


EXTRACTORS = {"pdf": ext_pdf, "hwp": ext_hwp, "hwpx": ext_hwpx, "docx": ext_docx, "pptx": ext_pptx, "xlsx": ext_xlsx,
              "xls": ext_xls, "doc": ext_doc, "txt": ext_text, "md": ext_text, "csv": ext_text}


def extract(path: str) -> tuple[list, dict]:
    ext = path.rsplit(".", 1)[-1].lower() if "." in os.path.basename(path) else ""
    if ext not in EXTRACTORS:
        raise RuntimeError(f"미지원 형식 .{ext}")
    return EXTRACTORS[ext](path)


# ---------------------------------------------------------------- 절 나누기
HEAD_RE = re.compile(
    r"^\s*(제\s?\d{1,3}\s?[장절편조]\b|\d{1,2}(\.\d{1,2}){0,2}\.?\s+\S|[IVX]{1,4}\.\s|[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]+\.?\s|"
    r"[가나다라마바사아자차카타파하]\.\s|Chapter\s?\d|CHAPTER\s?\d|[□■◆◇●◎▶]\s?\S|\d{1,2}\)\s|\[\s?\d{1,2}\s?\])")


def is_heading(line: str) -> bool:
    s = line.strip()
    if not (3 <= len(s) <= 48) or not HEAD_RE.match(s) or not re.search(r"[가-힣A-Za-z]", s):
        return False
    if re.search(r"[·…\.]{4,}\s*\d*\s*$", s):   # 목차 줄(점선 + 쪽수)
        return False
    if re.search(r"[.。]$|다$|요$|음$|함$|됨$|임$", s) and not re.match(r"^\d{1,2}(\.\d{1,2})*\.$", s):
        return False
    if len(re.findall(r"\d", s)) > max(6, len(s) // 2):   # 숫자만 나열된 표 행
        return False
    return True


def drop_running_lines(pages: list) -> list:
    """머리글·바닥글(쪽마다 반복되는 짧은 줄) 제거 — PDF 전용."""
    if len(pages) < 6:
        return pages
    from collections import Counter
    c = Counter()
    for pg in pages:
        seen = set()
        for ln in pg.splitlines():
            k = re.sub(r"\s+|\d+", "", ln.strip())
            if 2 <= len(k) <= 40 and k not in seen:
                seen.add(k); c[k] += 1
    th = max(4, int(len(pages) * 0.3))
    bad = {k for k, n in c.items() if n >= th}
    out = []
    for pg in pages:
        keep = []
        for ln in pg.splitlines():
            k = re.sub(r"\s+|\d+", "", ln.strip())
            if k in bad or re.fullmatch(r"\s*[-–—]?\s*\d{1,3}\s*[-–—]?\s*", ln):
                continue
            keep.append(ln)
        out.append("\n".join(keep))
    return out


def split_sections(pages: list, page_offset: int = 1, excluded: set | None = None) -> list:
    """pages(쪽 텍스트 목록) → [{'title','text','p0','p1'}] — 머리글 줄에서 자르되 MIN/MAX 길이를 지킨다. 쪽 번호는 1부터."""
    lines = []   # (page_no, line)
    for i, pg in enumerate(pages):
        pno = i + page_offset
        if excluded and pno in excluded:
            continue
        for ln in pg.splitlines():
            lines.append((pno, ln.rstrip()))
        lines.append((pno, ""))
    secs, cur, title, p0 = [], [], "", None

    def flush(next_title: str, next_p: int | None):
        nonlocal cur, title, p0
        txt = re.sub(r"\n{3,}", "\n\n", "\n".join(l for _, l in cur)).strip()
        if txt:
            secs.append({"title": title, "text": txt, "p0": p0, "p1": cur[-1][0]})
        cur, title, p0 = [], next_title, next_p

    size, last_p = 0, None
    for pno, ln in lines:
        if p0 is None:
            p0 = pno
        if last_p is not None and pno > last_p + 1 and size >= 300:   # 제외한 쪽을 건너뛴 자리에서는 끊는다
            flush("", pno); size = 0
        last_p = pno
        if is_heading(ln) and size >= MIN_SECTION:
            flush(ln.strip(), pno); size = 0
        elif size >= MAX_SECTION and ln.strip() == "":
            flush("", pno); size = 0
        cur.append((pno, ln)); size += len(ln) + 1
        if not title and is_heading(ln) and size < 200:
            title = ln.strip()
    flush("", None)
    # 제목이 비면 "(이어짐)" 표시
    prev = ""
    for s in secs:
        if not s["title"]:
            s["title"] = (prev + " (이어짐)") if prev else f"쪽 {s['p0']}~{s['p1']}"
        else:
            prev = re.sub(r"\s*\(이어짐\)$", "", s["title"])
    return secs


# ---------------------------------------------------------------- 분류·최신성
KIND_PAT = {
    "procedure": re.compile(r"클릭|버튼|메뉴|화면|입력|선택|로그인|경로|>|→|절차|순서|단계|체크리스트|신청서\s*작성|등록\s*방법", re.I),
    "case": re.compile(r"불인정|사례|Q\s*&\s*A|Q\.|질문|답변|O\s*/\s*X|가능\s*여부|예시|FAQ", re.I),
    "rule": re.compile(r"기준|한도|이내|이하|이상|전결|승인|금지|불가|원칙|규정|지침|의무|필수", re.I),
}


def guess_kind(text: str, ext: str) -> str:
    if ext in ("xlsx", "xls"):
        return "table"
    lines = [l for l in text.splitlines() if l.strip()]
    if lines and sum(1 for l in lines if l.count("|") >= 2) / len(lines) > 0.5:
        return "table"
    n = max(1, len(text) / 1000)
    sc = {k: len(p.findall(text)) / n for k, p in KIND_PAT.items()}
    best = max(sc, key=sc.get)
    return best if sc[best] >= 3 else "reference"


def doc_date_from(fname: str, meta: dict, page_date: str, m: dict) -> tuple[str, str]:
    if m.get("doc_date"):
        return m["doc_date"], "manifest"
    mm = re.search(r"(20\d{2})[.\-_]?(0[1-9]|1[0-2])[.\-_]?(0[1-9]|[12]\d|3[01])?", fname)
    if mm:
        return f"{mm.group(1)}-{mm.group(2)}", "filename"
    mm = re.search(r"(?<!\d)(2[0-3])(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])(?!\d)", fname)   # yymmdd
    if mm:
        return f"20{mm.group(1)}-{mm.group(2)}", "filename"
    mm = re.search(r"(?<!\d)(20\d{2})(?!\d)", fname)
    if mm:
        return mm.group(1), "filename"
    cd = meta.get("pdf_create") or ""
    mm = re.search(r"D:(20\d{2})(\d{2})", cd)
    if mm:
        return f"{mm.group(1)}-{mm.group(2)}", "pdfmeta"
    return (page_date or "")[:7], "page"


def years_old(doc_date: str) -> float:
    try:
        y, mo = int(doc_date[:4]), int(doc_date[5:7]) if len(doc_date) >= 7 else 6
        return (TODAY.year - y) + (TODAY.month - mo) / 12
    except Exception:
        return 0


def load_rule_changes() -> list:
    """`_shared/rule_changes.md` 의 표에서 (id, 항목, 옛 값, 현행 값, 시행일, 검색 패턴) 을 읽는다."""
    out = []
    if not os.path.exists(RULE_CHANGES):
        return out
    for ln in io.open(RULE_CHANGES, encoding="utf-8-sig"):
        if not ln.startswith("|") or ln.startswith("|--") or ln.startswith("| id") or "검색 패턴" in ln:
            continue
        cells = [c.strip() for c in re.split(r"(?<!\\)\|", ln.strip().strip("|"))]   # 패턴 안의 `\|` 는 표 구분자가 아님
        if len(cells) < 7 or not cells[6]:
            continue
        pat = cells[6].strip("`").replace("\\|", "|")
        try:
            out.append({"id": cells[0], "item": cells[1], "old": cells[2], "new": cells[3], "since": cells[4], "re": re.compile(pat)})
        except re.error:
            continue
    return out


# ---------------------------------------------------------------- 위키 페이지 겹침(superseded) 판정
def _norm(t: str) -> str:
    return re.sub(r"[\s\W_]+", "", t).lower()


def _shingles(t: str, k: int = 12, mod: int = 7) -> set:
    s = _norm(t)
    out = set()
    for i in range(0, max(0, len(s) - k + 1)):
        h = hash(s[i:i + k])
        if h % mod == 0:
            out.add(h)
    return out


def _norm_title(t: str) -> str:
    t = re.sub(r"^[\s\d.\-–—()\[\]가-힣]{0,4}?(?=[가-힣A-Za-z])", "", t.strip()) if re.match(r"^[\d.\-–—()\[\] ]+", t) else t
    t = re.sub(r"^(제\s?\d+\s?[장절편]|\d+(\.\d+)*\.?|[IVXⅠ-Ⅹ]+\.?|[□■◆●▶]|\[\d+\])\s*", "", t.strip())
    return _norm(t)


class WikiMatcher:
    def __init__(self, root: str, pages: list):
        self.pages = {p["id"]: p for p in pages}
        self.inv = {}      # shingle → set(page id)
        self.size = {}
        self.titles = {}
        for p in pages:
            fp = os.path.join(root, "pages", p["rel"].replace("/", os.sep)) if p.get("rel") else ""
            try:
                body = io.open(fp, encoding="utf-8-sig").read()
                m = re.match(r"---\n.*?\n---\n", body, re.S)
                body = body[m.end():] if m else body
            except Exception:
                body = ""
            sh = _shingles(body)
            self.size[p["id"]] = len(sh)
            for h in sh:
                self.inv.setdefault(h, set()).add(p["id"])
            nt = _norm_title(p.get("title") or "")
            if len(nt) >= 4:
                self.titles.setdefault(nt, []).append(p["id"])

    def match(self, title: str, text: str) -> tuple[str, float, str]:
        """(page id, containment, how) — 절 본문의 표본 조각 중 그 페이지에도 있는 비율. 제목이 같으면 문턱을 낮춘다."""
        sh = _shingles(text)
        if len(sh) < 12:
            return "", 0.0, ""
        cnt = {}
        for h in sh:
            for pid in self.inv.get(h, ()):
                cnt[pid] = cnt.get(pid, 0) + 1
        if not cnt:
            return "", 0.0, ""
        pid = max(cnt, key=cnt.get)
        c = cnt[pid] / len(sh)
        tset = set(self.titles.get(_norm_title(title), []))
        if c >= 0.45:
            return pid, round(c, 2), "본문 겹침"
        if tset:
            tp = max(tset, key=lambda x: cnt.get(x, 0))
            tc = cnt.get(tp, 0) / len(sh)
            if tc >= 0.15:
                return tp, round(tc, 2), "제목 같음+본문 일부"
        return "", round(c, 2), ""


# ---------------------------------------------------------------- 한 파일 처리
def yaml_s(s) -> str:
    return json.dumps("" if s is None else str(s), ensure_ascii=False)


def process_file(root: str, path: str, page: dict | None, manifest: dict, matcher: WikiMatcher | None, rules: list,
                 out_root: str, nested: str = "") -> list:
    """한 파일 → 절 목록(index 항목). page 는 위키 페이지(index.json 항목) 또는 None(docs/)."""
    fname = nested or os.path.basename(path)
    ext = fname.rsplit(".", 1)[-1].lower() if "." in fname else ""
    m = manifest_for(manifest, fname)
    page_id = page["id"] if page else "docs"
    page_path = page["path"] if page else "docs"
    page_url = page.get("url", "") if page else ""
    page_date = (page.get("updatedAt") or "")[:10] if page else ""
    base = {"page_id": page_id, "page_path": page_path, "page_url": page_url, "file": fname, "ext": ext}
    if m.get("skip"):
        return [dict(base, status="skip", note=m.get("note", "manifest skip"))]
    if ext in LIST_ONLY:
        return [dict(base, status="list_only", note="그림 파일 — 목록만")]
    if ext == "zip":
        # ZIP 안 파일은 **안쪽 경로 그대로** 식별(a/same.txt 와 b/same.txt 가 같은 출력 폴더를 써 앞 것이 지워지던 결함 — Codex 검토 v0.7.8),
        #   임시 파일도 항목마다 따로 폴더. 이름은 UTF-8 표시(flag 0x800)가 있으면 그대로, 없으면 한국 윈도우 압축(cp949)으로 읽어 본다.
        items = []
        try:
            with zipfile.ZipFile(path) as z:
                tmpd = tempfile.mkdtemp(prefix="kkwiki_zip_")
                try:
                    for k, zi in enumerate(z.infolist()):
                        if zi.is_dir():
                            continue
                        inner = zi.filename
                        if not (zi.flag_bits & 0x800):
                            try:
                                inner = inner.encode("cp437").decode("cp949")
                            except Exception:
                                pass
                        inner = inner.replace("\\", "/").lstrip("/")
                        iext = inner.rsplit(".", 1)[-1].lower() if "." in os.path.basename(inner) else ""
                        if iext in SUPPORTED and iext != "zip":
                            sub = os.path.join(tmpd, f"{k:04d}")
                            os.makedirs(sub, exist_ok=True)
                            dest = os.path.join(sub, safe_seg(os.path.basename(inner), 100))
                            open(dest, "wb").write(z.read(zi))
                            items += process_file(root, dest, page, manifest, matcher, rules, out_root, nested=f"{fname}::{inner}")
                        else:
                            items.append(dict(base, file=f"{fname}::{inner}", ext=iext, status="list_only", note="zip 안 파일 — 목록만"))
                finally:
                    shutil.rmtree(tmpd, ignore_errors=True)
        except Exception as e:
            items.append(dict(base, status="error", note=f"zip 열기 실패: {str(e)[:80]}"))
        return items
    if ext not in SUPPORTED:
        return [dict(base, status="list_only", note=f"미지원 형식 .{ext} — 목록만")]
    try:
        pages, meta = extract(path)
    except Exception as e:
        return [dict(base, status="error", note=f"추출 실패: {str(e)[:100]}")]
    full = "\n".join(pages)
    if ext == "pdf":
        g = garbled_ratio(full)
        if g > 0.3:
            return [dict(base, status="garbled", note=f"PDF 글자 깨짐(폰트 인코딩, 비율 {g:.2f}) — 같은 이름의 hwpx/docx 가 있으면 그쪽, 없으면 링크로 열어 확인")]
        nz = len(re.sub(r"\s", "", full))
        if nz < max(30, 10 * len(pages)):
            return [dict(base, status="no_text", note=f"이미지 PDF(글자 없음, {len(pages)}쪽) — 링크로 열어 확인")]
        pages = drop_running_lines(pages)
    pages = [scrub(p) for p in pages]
    if len(re.sub(r"<그림>|<표>|\s", "", "\n".join(pages))) < 80:
        return [dict(base, status="no_text", note="글자 거의 없음(그림·표 자리표시만) — 링크로 열어 확인")]
    excluded = set()
    for a, b in m.get("exclude_pages", []):
        excluded.update(range(a, b + 1))
    secs = split_sections(pages, 1, excluded) if ext in ("pdf", "pptx", "xlsx", "xls") else split_sections(["\n".join(pages)], 1, None)
    if not secs:
        return [dict(base, status="no_text", note="추출된 글자 없음")]
    doc_date, dds = doc_date_from(fname, meta, page_date, m)
    doc_title = m.get("title") or re.sub(r"\.[^.]+$", "", fname.split("::")[-1])
    issuer = m.get("issuer", "")
    poff = m.get("print_offset", 0)
    # 출력 폴더 = 파일 이름 **확장자까지**(같은 페이지의 신청서.hwp 와 신청서.docx 가 한 폴더를 써서 나중 것이 앞 것을 지웠다 — 2026-09-30)
    stem = safe_seg(fname.replace("::", "__").replace("/", "__"), 80)
    if "::" in fname:   # ZIP 안 파일 — 긴 경로가 80자에서 잘려 같아지지 않게 전체 경로의 짧은 지문을 붙인다
        stem = safe_seg(stem, 70) + "_" + hashlib.sha1(fname.encode("utf-8")).hexdigest()[:6]
    outdir = os.path.join(out_root, page_id, stem)
    if os.path.isdir(outdir):
        shutil.rmtree(outdir)
    os.makedirs(outdir, exist_ok=True)
    items = []
    n = len(secs)
    for i, s in enumerate(secs, 1):
        rule = {}
        for r in m.get("rules", []):
            a, b = r.get("pages", [0, 10 ** 9])
            if ext != "pdf" or (s["p0"] is not None and a <= s["p0"] <= b):
                rule = r; break
        kind = rule.get("kind") or m.get("kind") or guess_kind(s["text"], ext)
        sup, ov, how = ("", 0.0, "")
        if matcher and not m.get("no_auto_supersede"):
            sup, ov, how = matcher.match(s["title"], s["text"])
        if not sup and rule.get("superseded_by"):
            sup, how = rule["superseded_by"], "manifest"
        sup_title = matcher.pages[sup]["path"] if (matcher and sup in matcher.pages) else ""
        tier = rule.get("tier") or m.get("tier") or (3 if (sup or years_old(doc_date) >= OLD_YEARS) else 2)
        stale = [f"{r['item']}: 구 {r['old']} → 현행 {r['new']} ({r['since']})" for r in rules if r["re"].search(s["text"])]
        title = s["title"]
        pg = f"{s['p0']}-{s['p1']}" if ext == "pdf" and s["p0"] is not None else ""
        pgp = f"{s['p0'] + poff}-{s['p1'] + poff}" if pg and poff else ""
        tslug = safe_seg(re.sub(r'[\\/:*?"<>|]', "_", title), 40)
        fn = f"{i:03d}_{tslug}.md"
        rel = os.path.relpath(os.path.join(outdir, fn), root).replace(os.sep, "/")
        fm = ["---", f"title: {yaml_s(title)}", f"kind: {'attach' if page else 'doc'}", f"section_kind: {kind}", f"tier: {tier}",
              f"doc_title: {yaml_s(doc_title)}", f"doc_date: {yaml_s(doc_date)}", f"doc_date_source: {dds}", f"issuer: {yaml_s(issuer)}",
              f"source_page_id: {yaml_s(page_id)}", f"source_page: {yaml_s(page_path)}", f"source_url: {page_url}",
              f"file: {yaml_s(fname)}", f"file_ext: {ext}", f"pages: {yaml_s(pg)}", f"pages_print: {yaml_s(pgp)}", f"part: {i}/{n}",
              f"superseded_by: {yaml_s(sup)}", f"superseded_title: {yaml_s(sup_title)}", f"superseded_how: {yaml_s(how)}", f"overlap: {ov}",
              f"stale_values: {json.dumps(stale, ensure_ascii=False)}", f"note: {yaml_s(rule.get('note') or m.get('note') or '')}",
              f"extracted: {TODAY.isoformat()}", f"chars: {len(s['text'])}", "---", ""]
        io.open(os.path.join(outdir, fn), "w", encoding="utf-8").write("\n".join(fm) + f"# {title}\n\n{s['text']}\n")
        items.append(dict(base, status="ok", id=f"att:{page_id}:{stem}:{i:03d}", title=title, section_kind=kind, tier=tier,
                          doc_title=doc_title, doc_date=doc_date, issuer=issuer, pages=pg, pages_print=pgp, part=f"{i}/{n}",
                          superseded_by=sup, superseded_title=sup_title, superseded_how=how, overlap=ov, stale_values=stale,
                          note=rule.get("note") or m.get("note") or "", rel=rel, chars=len(s["text"])))
    return items


# ---------------------------------------------------------------- 명령
def run(root: str, force: bool, only: str, exts: set | None = None) -> None:
    pidx = os.path.join(root, "index.json")
    pages = json.load(io.open(pidx, encoding="utf-8")).get("pages", []) if os.path.exists(pidx) else []
    by_id = {p["id"]: p for p in pages}
    manifest = load_manifest(root)
    rules = load_rule_changes()
    log(f"[kk-wiki] 변경표 {len(rules)}항목, manifest 문서 규칙 {len(manifest.get('docs', []))}건")
    matcher = WikiMatcher(root, pages)
    out_root = os.path.join(root, TEXT_DIR)
    os.makedirs(out_root, exist_ok=True)
    prev = {}
    ip = os.path.join(root, INDEX_JSON)
    if os.path.exists(ip):   # --force 여도 대상이 아닌 파일(--only/--ext 밖)의 이전 결과는 지킨다
        try:
            for it in json.load(io.open(ip, encoding="utf-8")).get("sections", []):
                prev.setdefault((it["page_id"], it["file"].split("::")[0]), []).append(it)
        except Exception:
            prev = {}
    targets = []   # (path, page or None)
    adir = os.path.join(root, "attachments")
    if os.path.isdir(adir):
        for pid in sorted(os.listdir(adir)):
            for fn in sorted(os.listdir(os.path.join(adir, pid))):
                targets.append((os.path.join(adir, pid, fn), by_id.get(pid, {"id": pid, "path": "(index 에 없는 페이지)", "url": ""})))
    ddir = os.path.join(root, "docs")
    if os.path.isdir(ddir):
        for fn in sorted(os.listdir(ddir)):
            if os.path.isfile(os.path.join(ddir, fn)):
                targets.append((os.path.join(ddir, fn), None))
    sections, n_ok = [], 0
    stat = {}
    for path, page in targets:
        fn = os.path.basename(path)
        key = (page["id"] if page else "docs", fn)
        if only and only not in fn and only not in (page["path"] if page else "docs"):
            sections += prev.get(key, []); continue
        if exts and (fn.rsplit(".", 1)[-1].lower() if "." in fn else "") not in exts:
            sections += prev.get(key, []); continue
        old = prev.get(key)
        if old and not force and all(it.get("status") != "error" for it in old) and old[0].get("mtime") == int(os.path.getmtime(path)):
            sections += old; continue
        items = process_file(root, path, page, manifest, matcher, rules, out_root)
        mt = int(os.path.getmtime(path))
        for it in items:
            it["mtime"] = mt
        sections += items
        st = items[0]["status"] if items else "empty"
        stat[st] = stat.get(st, 0) + 1
        n_ok += sum(1 for it in items if it["status"] == "ok")
        tag = "+" if st == "ok" else "!"
        log(f"  {tag} {(page['path'] if page else 'docs')[:50]} / {fn[:50]} — {len([i for i in items if i['status'] == 'ok'])}절"
            + ("" if st == "ok" else f" [{st}] {items[0].get('note', '')[:70]}"))
    meta = {"built_at": datetime.datetime.now().isoformat(timespec="seconds"), "files": len(targets), "sections": len([s for s in sections if s["status"] == "ok"]),
            "sections_list": sections}
    io.open(ip, "w", encoding="utf-8").write(json.dumps({"built_at": meta["built_at"], "files": meta["files"], "sections": sections}, ensure_ascii=False, indent=0))
    write_index_md(root, sections)
    log(f"[kk-wiki] 추출: 파일 {len(targets)}, 절 {meta['sections']}, 상태 {stat} → {out_root}")


def write_index_md(root: str, sections: list) -> None:
    lines = [f"# 첨부·문서 절 목록 — {datetime.datetime.now():%Y-%m-%d %H:%M}", "",
             "tier 1 현행(위키 본문) / 2 첨부 / 3 구버전(새 판 있음·3년 초과). `superseded` = 같은 내용을 옮겨 적은 위키 페이지가 있음(그 페이지를 인용).", ""]
    cur = None
    for s in sections:
        k = (s["page_id"], s["file"].split("::")[0])
        if k != cur:
            cur = k
            lines.append(f"\n## {s['page_path']} / {s['file'].split('::')[0]}")
        if s["status"] != "ok":
            lines.append(f"- ✗ {s['file']} — {s['status']}: {s.get('note', '')}")
            continue
        sup = f" ← 대체: {s['superseded_title']}" if s.get("superseded_by") else ""
        st = f" ⚠구값 {len(s['stale_values'])}" if s.get("stale_values") else ""
        pg = f" p.{s['pages']}" if s.get("pages") else ""
        lines.append(f"- [{s['tier']}·{s['section_kind']}·{s['doc_date']}] {s['title']}{pg} ({s['chars']}자){sup}{st}  `{s['rel']}`")
    io.open(os.path.join(root, INDEX_MD), "w", encoding="utf-8").write("\n".join(lines) + "\n")


def status(root: str) -> None:
    ip = os.path.join(root, INDEX_JSON)
    if not os.path.exists(ip):
        log(f"[kk-wiki] 추출 결과 없음: {ip} — `run`"); return
    d = json.load(io.open(ip, encoding="utf-8"))
    secs = d.get("sections", [])
    files = {}
    for s in secs:
        files.setdefault((s["page_id"], s["file"].split("::")[0]), []).append(s)
    ok = [s for s in secs if s["status"] == "ok"]
    log(f"[kk-wiki] {d.get('built_at')} — 파일 {len(files)}, 절 {len(ok)} ({sum(s['chars'] for s in ok):,}자), "
        f"tier2 {sum(1 for s in ok if s['tier'] == 2)} / tier3 {sum(1 for s in ok if s['tier'] == 3)}, 대체됨 {sum(1 for s in ok if s['superseded_by'])}, 구값 표시 {sum(1 for s in ok if s['stale_values'])}")
    bad = [(k, v[0]) for k, v in files.items() if v[0]["status"] != "ok"]
    from collections import Counter
    log("  상태: " + str(Counter(v[0]["status"] for v in files.values())))
    for k, s in bad[:60]:
        log(f"  ✗ [{s['status']}] {s['page_path'][:45]} / {s['file'][:50]} — {s.get('note', '')[:60]}")


def restale(root: str) -> None:
    """규정 변경표(_shared/rule_changes.md)를 고친 뒤 '⚠ 구값' 표시만 다시 매긴다 — 다시 추출하지 않는다(몇 초).
    절 파일의 stale_values 줄과 attachments_index.json/.md 를 고친다."""
    ip = os.path.join(root, INDEX_JSON)
    if not os.path.exists(ip):
        log(f"[kk-wiki] 추출 결과 없음: {ip} — `run`"); return
    d = json.load(io.open(ip, encoding="utf-8"))
    rules = load_rule_changes()
    n_ok = n_chg = n_miss = 0
    for s in d.get("sections", []):
        if s.get("status") != "ok":
            continue
        fp = os.path.join(root, s["rel"].replace("/", os.sep))
        try:
            txt = io.open(fp, encoding="utf-8").read()
        except Exception:
            n_miss += 1; continue
        m = re.match(r"---\n.*?\n---\n", txt, re.S)
        body = txt[m.end():] if m else txt
        stale = [f"{r['item']}: 구 {r['old']} → 현행 {r['new']} ({r['since']})" for r in rules if r["re"].search(body)]
        n_ok += 1
        if stale != (s.get("stale_values") or []):
            s["stale_values"] = stale; n_chg += 1
            line = f"stale_values: {json.dumps(stale, ensure_ascii=False)}"
            io.open(fp, "w", encoding="utf-8").write(re.sub(r"^stale_values: .*$", lambda _m: line, txt, count=1, flags=re.M))
    io.open(ip, "w", encoding="utf-8").write(json.dumps(d, ensure_ascii=False, indent=0))
    write_index_md(root, d.get("sections", []))
    log(f"[kk-wiki] 구값 표시 다시: 변경표 {len(rules)}항목, 절 {n_ok}개 중 {n_chg}개 바뀜"
        + (f", 파일 없음 {n_miss}(run 으로 다시 추출)" if n_miss else "") + f" — 지금 구값 표시 {sum(1 for s in d.get('sections', []) if s.get('stale_values'))}절")


def show(root: str, pat: str) -> None:
    d = json.load(io.open(os.path.join(root, INDEX_JSON), encoding="utf-8"))
    for s in d.get("sections", []):
        if pat.lower() in s["file"].lower() or pat.lower() in s["page_path"].lower():
            if s["status"] != "ok":
                log(f"✗ {s['file']} [{s['status']}] {s.get('note', '')}"); continue
            sup = f" ← {s['superseded_title']} ({s['superseded_how']} {s['overlap']})" if s["superseded_by"] else ""
            log(f"{s['part']:>7} [{s['tier']}·{s['section_kind']}] p.{s['pages'] or '-'} {s['title'][:40]} ({s['chars']}자){sup}"
                + (f"  ⚠{s['stale_values']}" if s["stale_values"] else ""))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", choices=["run", "status", "show", "restale"])
    ap.add_argument("args", nargs="*")
    ap.add_argument("--root")
    ap.add_argument("--force", action="store_true", help="이미 추출한 파일도 다시")
    ap.add_argument("--only", default="", help="파일 이름·페이지 경로에 이 문자열이 든 것만 (나머지는 이전 결과 유지)")
    ap.add_argument("--ext", default="", help="이 확장자만 (쉼표 구분, 예 hwp,zip) — 나머지는 이전 결과 유지")
    a = ap.parse_args(argv)
    root = snapshot_root(a.root)
    if a.cmd == "run":
        run(root, a.force, a.only, {e.strip().lower().lstrip(".") for e in a.ext.split(",") if e.strip()} or None)
    elif a.cmd == "status":
        status(root)
    elif a.cmd == "restale":
        restale(root)
    else:
        if not a.args:
            raise SystemExit("show <파일 문자열>")
        show(root, a.args[0])


if __name__ == "__main__":
    main()
