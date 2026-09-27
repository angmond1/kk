# -*- coding: utf-8 -*-
r"""
kk-meeting 회의록 엑셀 헬퍼.

표준(2026-09-24~): `{root}\{yymm}_회의록.xlsx` — 한 폴더, 월별 1파일(yymm = 지급신청 처리 연월). 그 달 처리 건은 모두 같은 파일에 행 추가.
주된 목적 = 이전 회의 주제·내용과의 중복 방지 기록(지급신청에 첨부하지 않음). 새 회의록 전 all_titles() 로 전부 스캔.

9컬럼 형식:
  순번 | 사용일자 | 금액 | 장소(거래처) | 처리계정 | 내부참석자 | 외부참석자 | 외부참석자 소속 | 회의목적

회의목적 셀에는 "제목\n상세내용" 형태로 작성 (10만원 ↑ 상세 필수).

사용 예:
    from meeting_log_xlsx import open_or_create, append_row, read_log
    path = open_or_create("2606")            # {root}\2606_회의록.xlsx
    append_row(path, {
        "date_text": "4월 30일 13:00~14:30",
        "amount": 150000,
        "place": "○○식당",
        "acccd": "2E11111",
        "int_members": "김키키",
        "ext_members": "이키키, 박키키, 최키키",
        "ext_org": "○○대학교",
        "title": "연구 진행상황 논의",
        "content": "1. 연구 진행상황 및 향후 계획 공유\n - ...",
    })
    rows = read_log(path)   # list[dict]
    past = all_titles()      # 중복 방지: 폴더 내 모든 회의록의 제목·내용
"""
from __future__ import annotations
import os
import re
import sys
# Windows 한국어(cp949) 콘솔·파이프에서도 한글·기호가 깨지거나 멈추지 않게 출력은 UTF-8 로 (모듈로 불러 써도 적용)
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
from typing import Optional

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None
_NEED = "openpyxl 이 없습니다: python -m pip install openpyxl  (macOS/Linux: python3 -m pip install --user openpyxl)"

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
HEADERS = ["순번", "사용일자", "금액", "장소(거래처)", "처리계정",
           "내부참석자", "외부참석자", "외부참석자 소속", "회의목적"]
WIDTHS = [6, 20, 11, 18, 10, 12, 26, 18, 55]

def _kiki_root() -> str:
    """kiki 작업 폴더 — 환경변수 KIKI_ROOT → kiki.config.json(claude/codex) 의 kiki_root → 기본(C:\\kiki / ~/kiki)."""
    r = os.environ.get("KIKI_ROOT", "").strip()
    if r:
        return os.path.expanduser(r)
    for cfg in ("~/.claude/kiki/kiki.config.json", "~/.codex/kiki/kiki.config.json"):
        p = os.path.expanduser(cfg)
        if os.path.exists(p):
            try:
                import json
                r = (json.load(open(p, encoding="utf-8-sig")).get("kiki_root") or "").strip()
            except Exception:
                r = ""
            if r:
                return os.path.expanduser(r)
    return r"C:\kiki" if os.name == "nt" else os.path.expanduser("~/kiki")


def default_root() -> str:
    """회의록 엑셀 폴더 — kk-meeting.config.json 의 log_root({kiki_root} 치환) → 없으면 {kiki_root}/meeting/meeting_log."""
    for cfg in ("~/.claude/kiki/kk-meeting.config.json", "~/.codex/kiki/kk-meeting.config.json"):
        p = os.path.expanduser(cfg)
        if os.path.exists(p):
            try:
                import json
                lr = (json.load(open(p, encoding="utf-8-sig")).get("log_root") or "").strip()
            except Exception:
                lr = ""
            if lr:
                return os.path.expanduser(lr.replace("{kiki_root}", _kiki_root()).replace("/dining/", "/meeting/").replace("\\dining\\", "\\meeting\\"))
    return os.path.join(_kiki_root(), "meeting", "meeting_log")


# 하위 호환 이름 — 모듈을 불러온 시점의 기본 폴더 (함수 인자 root 가 우선)
DEFAULT_ROOT = default_root()


def expected_path(yymm: str, root: Optional[str] = None) -> str:
    r"""`{root}\{yymm}_회의록.xlsx` 경로 반환 (한 폴더, 월별 1파일)."""
    root = root or DEFAULT_ROOT
    return os.path.join(root, f"{yymm}_회의록.xlsx")


def open_or_create(yymm: str, root: Optional[str] = None) -> str:
    """엑셀 파일을 열거나(있으면) 9컬럼 양식으로 생성(없으면). 경로 반환."""
    if openpyxl is None:
        raise RuntimeError(_NEED)
    fp = expected_path(yymm, root)
    folder = os.path.dirname(fp)
    os.makedirs(folder, exist_ok=True)
    if os.path.exists(fp):
        return fp

    wb = Workbook()
    ws = wb.active
    ws.title = "회의록"
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(fp)
    return fp


def _to_amount(v) -> int:
    """323000 / '323,000' / '45,000
11,000'(같은 날 식당+카페 두 금액) → 정수 합계. 숫자 없으면 0."""
    if isinstance(v, (int, float)):
        return int(v)
    text = re.sub(r"\([^)]*\)", " ", str(v or ""))          # 괄호 안 내역(식대 30,000 + 음료 15,000)은 합계에 안 더한다
    nums = re.findall(r"-?\d[\d,]*", text)
    return sum(int(n.replace(",", "")) for n in nums) if nums else 0


def _next_seq(ws) -> int:
    """현재 시트의 다음 순번."""
    seq = 0
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        v = row[0]
        if isinstance(v, (int, float)):
            seq = max(seq, int(v))
    return seq + 1


def append_row(path: str, data: dict) -> int:
    """1건 추가. data 키:
        date_text   사용일자 (예 '4월 30일 13:00~14:30')
        amount      금액 (int)
        place       장소(거래처)
        acccd       처리계정 (예 '2E11111')
        int_members 내부참석자 (성명, 콤마 구분)
        ext_members 외부참석자 (성명, 콤마 구분)
        ext_org     외부참석자 소속
        title       회의제목
        content     회의내용 (10만원↑ 상세)
    return: 부여된 순번
    """
    if openpyxl is None:
        raise RuntimeError(_NEED)
    wb = openpyxl.load_workbook(path)
    ws = wb["회의록"] if "회의록" in wb.sheetnames else wb.active
    seq = _next_seq(ws)

    def _clean(v):                                                 # 엑셀이 거부하는 제어문자(PDF 복사 시 섞이는 수직탭·쪽나눔 등) 제거
        return _ILLEGAL.sub(" ", v) if isinstance(v, str) else v
    data = {k: _clean(v) for k, v in (data or {}).items()}
    title = data.get("title", "")
    content = data.get("content", "")
    purpose_cell = (f"{title}\n{content}" if content else title)

    ws.append([
        seq,
        data.get("date_text", ""),
        _to_amount(data.get("amount", 0)),
        data.get("place", ""),
        data.get("acccd", ""),
        data.get("int_members", ""),
        data.get("ext_members", ""),
        data.get("ext_org", ""),
        purpose_cell,
    ])
    r = ws.max_row
    # 셀 포맷
    ws.cell(row=r, column=3).number_format = "#,##0"
    ws.cell(row=r, column=7).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r, column=8).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r, column=9).alignment = Alignment(wrap_text=True, vertical="top")
    try:
        wb.save(path)
    except PermissionError:
        raise RuntimeError(f"엑셀 파일이 열려 있어 저장하지 못했습니다 — Excel 에서 닫고 다시: {path}")
    return seq


def read_log(path: str) -> list[dict]:
    """엑셀 회의록 전체 행을 dict 리스트로 반환 (Claude 가 fam_0704 작성 시 조회)."""
    if openpyxl is None:
        raise RuntimeError(_NEED)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["회의록"] if "회의록" in wb.sheetnames else wb.active
    out = []
    for row in ws.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True):
        if not row or row[0] is None:
            continue
        purpose = row[8] or ""
        # 회의목적 셀 분리 (제목 첫 줄 / 상세 이후)
        if "\n" in (purpose or ""):
            title, content = purpose.split("\n", 1)
        else:
            title, content = purpose, ""
        out.append({
            "seq": row[0],
            "date_text": row[1],
            "amount": row[2],
            "place": row[3],
            "acccd": row[4],
            "int_members": row[5],
            "ext_members": row[6],
            "ext_org": row[7],
            "title": title,
            "content": content,
        })
    return out


def find_duplicate(path: str, data: dict) -> Optional[dict]:
    """같은 회의(사용일자·장소·계정·금액이 모두 같은 행)가 그 파일에 이미 있으면 그 행(dict), 없으면 None."""
    if not os.path.exists(path):
        return None
    key = (str(data.get("date_text", "")).strip(), str(data.get("place", "")).strip(), str(data.get("acccd", "")).strip(), _to_amount(data.get("amount", 0)))
    for r in read_log(path):
        if (str(r["date_text"] or "").strip(), str(r["place"] or "").strip(), str(r["acccd"] or "").strip(), _to_amount(r["amount"])) == key:
            return r
    return None


def all_titles(root: Optional[str] = None, skipped: Optional[list] = None) -> list[dict]:
    """중복 방지용 — 폴더의 모든 `*_회의록.xlsx` 를 읽어
    [{file, date_text, amount, place, acccd, title, content}] 를 반환한다.
    새 회의록을 쓰기 전 사용자가 준 주제를 title 들과 비교(같거나 유사하면 조정 제안)."""
    root = root or DEFAULT_ROOT
    out: list[dict] = []
    if not os.path.isdir(root):
        return out
    files = [fn for fn in sorted(os.listdir(root)) if fn.endswith("_회의록.xlsx") and not fn.startswith("~$")]
    if files and openpyxl is None:                     # 회의록이 아직 없으면 openpyxl 없이도 0건(새 PC 첫 회의록 — 2026-09-27 시험에서 발견)
        raise RuntimeError(_NEED)
    for fn in files:
        try:
            wb = openpyxl.load_workbook(os.path.join(root, fn), data_only=True)
            ws = wb["회의록"] if "회의록" in wb.sheetnames else wb.active
        except Exception as e:
            if skipped is not None:
                skipped.append(f"{fn}({type(e).__name__})")
            continue
        for r in ws.iter_rows(min_row=2, values_only=True):
            r = list(r) + [None] * 9
            if not any(r[:9]):
                continue
            purpose = str(r[8] or "")
            title, _, content = purpose.partition("\n")
            out.append({"file": fn, "date_text": str(r[1] or ""), "amount": r[2], "place": str(r[3] or ""),
                        "acccd": str(r[4] or ""), "title": title.strip(), "content": content.strip()})
    return out


if __name__ == "__main__":
    # 명령줄 (Claude 가 한글이 든 파이썬 한 줄을 셸에 넣지 않게 — 2026-09-27):
    #   python meeting_log_xlsx.py append <row.json> [--yymm YYMM] [--root DIR] [--dup-ok]  → 행 추가, "순번 N | 파일" + [요약]
    #   python meeting_log_xlsx.py titles [--root DIR] [--full]                           → 과거 회의 한 줄씩(중복 검사용). --full 이면 내용까지 JSON
    #   python meeting_log_xlsx.py path <YYMM> [--root DIR]                                → 그 달 파일 경로(없으면 생성)
    import json as _json, time as _time
    _USAGE = ("사용: python meeting_log_xlsx.py append <row.json> [--yymm YYMM] [--root DIR] [--dup-ok]\n"
              "      python meeting_log_xlsx.py titles [--root DIR] [--full]\n"
              "      python meeting_log_xlsx.py path <YYMM> [--root DIR]")
    _a = sys.argv[1:]
    def _opt(name):
        return _a[_a.index(name) + 1] if name in _a and _a.index(name) + 1 < len(_a) else None
    _root = _opt("--root")
    try:
        if _a and _a[0] == "append" and len(_a) >= 2:
            _data = _json.load(open(_a[1], encoding="utf-8-sig"))
            if not isinstance(_data, dict):
                print("ERR row.json 은 객체 하나({date_text, amount, place, acccd, int_members, ext_members, ext_org, title, content})여야 합니다"); sys.exit(1)
            _miss = [k for k in ("date_text", "amount", "place", "acccd", "title") if not str(_data.get(k) or "").strip() or (k == "amount" and not _to_amount(_data.get(k)))]
            if _miss:
                print("ERR 필수 항목이 비어 기록하지 않았습니다:", ", ".join(_miss)); sys.exit(1)
            _p = open_or_create(_opt("--yymm") or _time.strftime("%y%m"), _root)
            _dup = find_duplicate(_p, _data)
            if _dup and "--dup-ok" not in _a:
                print(f"ERR 같은 회의(사용일자·장소·계정·금액)가 이미 순번 {_dup['seq']} 으로 기록돼 있어 다시 쓰지 않았습니다 — 정말 별개 회의면 --dup-ok")
                sys.exit(1)
            _seq = append_row(_p, _data)
            print("순번", _seq, "|", _p)
            print(f"[요약] 추가 1건 — 이 파일 {len(read_log(_p))}건(직전 순번 {_seq - 1})")
        elif _a and _a[0] == "titles":
            _skip = []
            _rows = all_titles(_root, _skip)
            if "--full" in _a:
                print(_json.dumps(_rows, ensure_ascii=False, indent=1))
            else:
                print(f"[회의록 {len(_rows)}건 | 파일 {len({r['file'] for r in _rows})}개 | {_root or DEFAULT_ROOT}] 월 | 사용일자 | 계정 | 제목")
                for _r in _rows:
                    print(f"{_r['file'][:4]} | {_r['date_text'][:18]} | {_r['acccd']} | {_r['title'][:70]}")
            if _skip:
                print("⚠ 읽지 못한 회의록 파일(열려 있거나 손상 — 중복 검사에서 빠짐):", ", ".join(_skip))
        elif _a and _a[0] == "path" and len(_a) >= 2:
            print(open_or_create(_a[1], _root))
        else:
            print(_USAGE)
            sys.exit(1)
    except RuntimeError as e:
        print("ERR", e)
        sys.exit(1)
