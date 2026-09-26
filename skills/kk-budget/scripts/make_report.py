# -*- coding: utf-8 -*-
"""
kk-budget 예산 리포트 렌더러.
입력: fetch 수집 JSON (snapshot_date / track_categories / projects) — projects 는 {acccd:{…}} 사전 또는 [{acccd,…}] 목록 둘 다 됨.
  과제 = {name, pi, role, direct{A,D}, categories:{표시명:{A,exec,pendingDone,pendingProg,D}}} (kkBudget.queryBudgetTable 반환 그대로; 없는 비목은 키를 빼야 '-' 로 표시)
출력: 과제 행 x 카테고리(총액/잔액) + (한 칸 띄우고) 직접비(잔액/총액) 엑셀.
사용: python make_report.py <input.json> <output.xlsx>
credential-free. 개인 식별자/경로 하드코딩 없음(전부 인자/JSON).
"""
import json
import sys
# Windows 한국어(cp949) 콘솔·파이프에서도 한글·기호가 깨지거나 멈추지 않게 출력은 UTF-8 로 (모듈로 불러 써도 적용)
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
import os
from pathlib import Path
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("[kk-budget] openpyxl 이 없습니다: python -m pip install openpyxl  (macOS/Linux 는 python3 -m pip install --user openpyxl)")

THIN = Side(style="thin", color="999999")
LIGHT = Side(style="thin", color="CCCCCC")


def h_main(c):
    c.font = Font(bold=True, color="FFFFFF", size=11)
    c.fill = PatternFill("solid", fgColor="305496")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = Border(THIN, THIN, THIN, THIN)


def h_sub(c, color="D9E1F2"):
    c.font = Font(bold=True, size=10)
    c.fill = PatternFill("solid", fgColor=color)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = Border(THIN, THIN, THIN, THIN)


def money(c, bold=False, muted=False):
    c.number_format = '#,##0'
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.border = Border(LIGHT, LIGHT, LIGHT, LIGHT)
    c.font = Font(bold=bold, size=11, color="808080" if muted else "000000")


def txt(c, bold=False, align="left"):
    c.font = Font(bold=bold, size=11)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    c.border = Border(LIGHT, LIGHT, LIGHT, LIGHT)


def dash(ws, row, col, label="-"):
    for k in range(2):
        c = ws.cell(row=row, column=col + k, value=label)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = Font(color="BBBBBB")
        c.border = Border(LIGHT, LIGHT, LIGHT, LIGHT)


def main(json_path, out_path):
    out_path = os.path.expanduser(str(out_path))
    if "{" in out_path or "}" in out_path:
        sys.exit("[kk-budget] 출력 경로에 치환 안 된 자리표시({kiki_root} 등)가 있습니다: " + out_path)
    with open(os.path.expanduser(str(json_path)), encoding="utf-8-sig") as f:   # PowerShell 이 BOM 을 붙여 저장해도 읽힘
        snap = json.load(f)
    cats = snap["track_categories"]
    projs = snap.get("projects") or {}
    if isinstance(projs, list):                        # [{acccd,…}] 목록도 허용 (queryProjects·kiki.config 형식)
        projs = {str(pj.get("acccd", "")): pj for pj in projs if isinstance(pj, dict) and pj.get("acccd")}
    ds = snap["snapshot_date"]

    wb = Workbook()
    ws = wb.active
    ws.title = "예산 현황"
    ws["A1"] = f"과제별 예산 현황 — 기준일 {ds}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.row_dimensions[1].height = 24

    r1, r2 = 3, 4
    for i, name in enumerate(["과제번호", "과제명", "PI/역할"], start=1):
        c = ws.cell(row=r1, column=i, value=name)
        h_main(c)
        ws.merge_cells(start_row=r1, start_column=i, end_row=r2, end_column=i)

    col = 4
    for cat in cats:
        m = ws.cell(row=r1, column=col, value=cat)
        h_main(m)
        ws.merge_cells(start_row=r1, start_column=col, end_row=r1, end_column=col + 1)
        h_sub(ws.cell(row=r2, column=col, value="총액"))
        h_sub(ws.cell(row=r2, column=col + 1, value="잔액"), color="FCE4D6")
        col += 2

    # 학생인건비(마지막 카테고리)와 한 칸 띄우고 직접비(잔액/총액)
    blank_col = col          # 구분용 빈 열
    direct_col = col + 1     # 직접비 시작 열
    dm = ws.cell(row=r1, column=direct_col, value="직접비")
    h_main(dm)
    ws.merge_cells(start_row=r1, start_column=direct_col, end_row=r1, end_column=direct_col + 1)
    h_sub(ws.cell(row=r2, column=direct_col, value="잔액"), color="FCE4D6")
    h_sub(ws.cell(row=r2, column=direct_col + 1, value="총액"))

    ws.row_dimensions[r1].height = 22
    ws.row_dimensions[r2].height = 20

    row = r2 + 1
    for code, p in projs.items():
        txt(ws.cell(row=row, column=1, value=code), bold=True, align="center")
        txt(ws.cell(row=row, column=2, value=p.get("name", "")), bold=True)
        txt(ws.cell(row=row, column=3, value=f'{p.get("pi","")}/{p.get("role","")}'), align="center")
        col = 4
        for cat in cats:
            data = p.get("categories", {}).get(cat)
            if not data:
                dash(ws, row, col)
            else:
                money(ws.cell(row=row, column=col, value=data.get("A", 0)), muted=True)
                money(ws.cell(row=row, column=col + 1, value=data.get("D", 0)), bold=True)
            col += 2
        # 직접비 (빈 열 건너뛰고 잔액/총액)
        direct = p.get("direct")
        if direct:
            money(ws.cell(row=row, column=direct_col, value=direct.get("D", 0)), bold=True)
            money(ws.cell(row=row, column=direct_col + 1, value=direct.get("A", 0)), muted=True)
        ws.row_dimensions[row].height = 22
        row += 1

    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 12
    col = 4
    for _ in cats:
        ws.column_dimensions[get_column_letter(col)].width = 14
        ws.column_dimensions[get_column_letter(col + 1)].width = 14
        col += 2
    ws.column_dimensions[get_column_letter(blank_col)].width = 3      # 구분 빈 열(좁게)
    ws.column_dimensions[get_column_letter(direct_col)].width = 15
    ws.column_dimensions[get_column_letter(direct_col + 1)].width = 15
    ws.freeze_panes = ws.cell(row=r2 + 1, column=4)

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"[OK] 저장: {out} (과제 {len(projs)}, 비목 {len(cats)}, 기준일 {ds})")
    return out




def _kiki_root() -> str:
    r = os.environ.get("KIKI_ROOT", "").strip()
    if r:
        return os.path.expanduser(r)
    for cfg in ("~/.claude/kiki/kiki.config.json", "~/.codex/kiki/kiki.config.json"):
        p = os.path.expanduser(cfg)
        if os.path.exists(p):
            try:
                r = (json.load(open(p, encoding="utf-8-sig")).get("kiki_root") or "").strip()
            except Exception:
                r = ""
            if r:
                return os.path.expanduser(r)
    return "C:\\kiki" if os.name == "nt" else os.path.expanduser("~/kiki")


def _downloads_dir() -> str:
    d = os.environ.get("KIKI_DOWNLOADS", "").strip()
    if d:
        return os.path.expanduser(d)
    if os.name == "nt":
        try:
            import winreg
            k = winreg.OpenKey(winreg.HKEY_CURRENT_USER, r"Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders")
            v, _ = winreg.QueryValueEx(k, "{374DE290-123F-4565-9164-39C4925E467B}")
            return os.path.expandvars(v)
        except Exception:
            pass
    return os.path.join(os.path.expanduser("~"), "Downloads")


def _latest_download(pattern: str, max_age_h: float = 24.0) -> str:
    import glob, time
    cands = [p for p in glob.glob(os.path.join(_downloads_dir(), pattern)) if os.path.isfile(p) and time.time() - os.path.getmtime(p) <= max_age_h * 3600]
    return max(cands, key=os.path.getmtime) if cands else ""


def cli(argv: list) -> int:
    """사용: python make_report.py <input.json> <output.xlsx>
       python make_report.py --from-downloads [--out <output.xlsx>] [--save-snapshot [<dir>]]
         --from-downloads : 다운로드 폴더의 최근 kiki_budget_*.json(브라우저 코어 downloadSnapshot 결과)
         --out            : 기본 {kiki_root}/budget/yymmdd.xlsx (yymmdd = 스냅샷 날짜)
         --save-snapshot  : 입력 JSON 을 <dir>/yymmdd.json 으로 보관 (기본 ~/.claude/kiki/kk-budget/data)"""
    if not argv or argv[0] in ("-h", "--help"):
        print(cli.__doc__); return 0 if argv else 1
    import shutil
    src = out = None; save = None
    pos = []
    i = 0
    while i < len(argv):
        a = argv[i]
        if a == "--from-downloads":
            src = _latest_download("kiki_budget_*.json")
            if not src:
                print("[kk-budget] 다운로드 폴더(" + _downloads_dir() + ")에 24시간 내 kiki_budget_*.json 이 없습니다 — 브라우저 코어 downloadSnapshot() 먼저"); return 1
        elif a == "--out" and i + 1 < len(argv):
            out = argv[i + 1]; i += 1
        elif a == "--save-snapshot":
            save = argv[i + 1] if i + 1 < len(argv) and not argv[i + 1].startswith("--") else os.path.expanduser("~/.claude/kiki/kk-budget/data")
            if save == (argv[i + 1] if i + 1 < len(argv) else None):
                i += 1
        else:
            pos.append(a)
        i += 1
    if src is None and pos:
        src = pos.pop(0)
    if out is None and pos:
        out = pos.pop(0)
    if not src:
        print(cli.__doc__); return 1
    try:
        with open(os.path.expanduser(src), encoding="utf-8-sig") as f:
            snap = json.load(f)
    except Exception as e:
        print("[kk-budget] 입력 JSON 을 읽지 못함:", e); return 1
    ymd = (snap.get("snapshot_date") or "").replace("-", "")[2:] or __import__("time").strftime("%y%m%d")
    if not out:
        out = os.path.join(_kiki_root(), "budget", ymd + ".xlsx")
    if save:
        os.makedirs(os.path.expanduser(save), exist_ok=True)
        dst = os.path.join(os.path.expanduser(save), ymd + ".json")
        shutil.copy2(os.path.expanduser(src), dst); print("[kk-budget] 스냅샷 보관:", dst)
    main(src, out)
    return 0


if __name__ == "__main__":
    sys.exit(cli(sys.argv[1:]))
