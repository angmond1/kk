# -*- coding: utf-8 -*-
"""kk-wiki 첨부 추출(wiki_extract.py)·한 번에 모으기(wiki_ask.py) 회귀 시험 — 네트워크·토큰·실데이터 없이 임시 폴더에서만.
  python tools/selftest_wiki_extract.py [스크립트 폴더]   (기본 skills/kk-wiki/scripts — 옛 판과 비교할 때 다른 폴더를 줄 수 있다)
Codex 검토(v0.7.8)에서 재현한 결함:
  ① ZIP 안 a/same.txt·b/same.txt 가 같은 출력 폴더를 써 앞 파일 내용이 사라짐
  ② DOCX 에서 문단 뒤 표가 문서 끝으로 밀려 순서가 틀림
  ③ wiki_ask 가 긴 표를 행 가운데서 자르면서 '표 통째' 라고 안내함
  + 토큰이 없을 때 최신 확인 대신 붙여 넣을 브라우저 호출을 준다
"""
import io, json, os, shutil, subprocess, sys, tempfile, zipfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCR = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.join(REPO, "skills", "kk-wiki", "scripts")
PY = sys.executable
n = fail = 0


def chk(name, cond, detail=""):
    global n, fail
    n += 1
    if cond:
        print("PASS " + name)
    else:
        fail += 1
        print("FAIL " + name + ("  ← " + str(detail)[:600] if detail else ""))


def run(args, root, env_extra=None):
    env = dict(os.environ, PYTHONIOENCODING="utf-8", KIKI_ROOT=os.path.dirname(root), DOORAY_TOKEN="")
    env.update(env_extra or {})
    r = subprocess.run([PY] + args + ["--root", root], capture_output=True, text=True, encoding="utf-8", env=env, timeout=300)
    return r.returncode, (r.stdout or "") + (r.stderr or "")


T = tempfile.mkdtemp(prefix="kkwiki_selftest_")
try:
    H = os.path.join(T, "home"); os.makedirs(H)
    os.environ.update(HOME=H, USERPROFILE=H)   # 실제 ~/.claude/kiki 의 토큰·설정을 읽지 않게
    root = os.path.join(T, "kiki", "wiki")
    os.makedirs(os.path.join(root, "pages"))
    page = {"id": "1001", "title": "시험 쪽", "path": "시험 쪽", "rel": "시험 쪽.md", "url": "https://example.invalid/wiki/1001",
            "updatedAt": "2026-09-01T10:00:00+09:00", "version": 3, "len": 10, "files": ["묶음.zip", "순서.docx"]}
    json.dump({"space_id": "S", "home_page_id": "1001", "built_at": "2026-09-30", "count": 1, "pages": [page]},
              io.open(os.path.join(root, "index.json"), "w", encoding="utf-8"), ensure_ascii=False)
    io.open(os.path.join(root, "pages", "시험 쪽.md"), "w", encoding="utf-8").write("---\ntitle: \"시험 쪽\"\n---\n# 시험 쪽\n\n첨부 확인\n")
    att = os.path.join(root, "attachments", "1001"); os.makedirs(att)
    # ① 같은 이름, 다른 폴더
    with zipfile.ZipFile(os.path.join(att, "묶음.zip"), "w") as z:
        z.writestr("a/same.txt", "앞폴더 내용 알파입니다. 같은 이름 파일이 둘 있어도 따로 남아야 합니다. " * 4)
        z.writestr("b/same.txt", "뒤폴더 내용 베타입니다. 같은 이름 파일이 둘 있어도 따로 남아야 합니다. " * 4)
    # ② 문단 → 표 → 문단
    try:
        import docx
        d = docx.Document()
        d.add_paragraph("첫 문단 가나다 — 문단과 표가 섞인 문서는 문서 순서대로 읽어야 합니다. 이 문단이 맨 앞입니다.")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "항목", "값", "표칸 라마바", "42"
        d.add_paragraph("끝 문단 사아자 — 표 뒤에 오는 문단입니다. 이 문단이 맨 끝이어야 합니다.")
        d.save(os.path.join(att, "순서.docx"))
        have_docx = True
    except Exception as e:
        have_docx = False
        print("SKIP docx 시험 — python-docx 없음: " + str(e)[:60])
    # 이름은 같고 확장자만 다른 두 첨부(신청서.hwp·신청서.docx 같은 경우) — 출력 폴더가 겹쳐 앞 것이 지워지면 안 됨
    io.open(os.path.join(att, "같은이름.txt"), "w", encoding="utf-8").write("글 파일 쪽 내용 감마입니다. 확장자만 다른 파일이 같은 페이지에 있어도 둘 다 남아야 합니다. " * 3)
    io.open(os.path.join(att, "같은이름.md"), "w", encoding="utf-8").write("마크다운 쪽 내용 델타입니다. 확장자만 다른 파일이 같은 페이지에 있어도 둘 다 남아야 합니다. " * 3)
    # 구값 표시: 피험자 사례비의 옛 값(12만 5천원)은 표시, 자문료의 과세최저한 12만 5천원(현행)은 표시하지 않음
    io.open(os.path.join(att, "피험자.txt"), "w", encoding="utf-8").write("임상실험 피험자 사례비는 12만 5천원 초과 지급 시 소득세를 공제한다. 생명윤리심의 증빙과 피험자 동의서를 붙인다. " * 2)
    io.open(os.path.join(att, "자문료.txt"), "w", encoding="utf-8").write("국내 전문가 자문료는 과세최저한 12만 5천원 초과 시 기타소득세를 원천징수한다. 이하라도 원천세 정보는 입력한다. " * 2)
    rc, out = run([os.path.join(SCR, "wiki_extract.py"), "run"], root)
    idx = json.load(io.open(os.path.join(root, "attachments_index.json"), encoding="utf-8")) if os.path.exists(os.path.join(root, "attachments_index.json")) else {"sections": []}
    secs = [s for s in idx.get("sections", []) if s.get("status") == "ok"]

    def text_of(pred):
        return "\n".join(io.open(os.path.join(root, s["rel"].replace("/", os.sep)), encoding="utf-8").read() for s in secs if pred(s))

    zt = text_of(lambda s: "묶음.zip::" in s["file"])
    chk("ZIP 안 같은 이름 두 파일(a/same.txt·b/same.txt)이 둘 다 남음", "알파" in zt and "베타" in zt, out[-400:] + " // " + json.dumps([s["file"] for s in secs], ensure_ascii=False))
    st = text_of(lambda s: s["file"] in ("같은이름.txt", "같은이름.md"))
    chk("이름은 같고 확장자만 다른 두 첨부가 둘 다 남음(출력 폴더에 확장자)", "감마" in st and "델타" in st and all(os.path.exists(os.path.join(root, s["rel"].replace("/", os.sep))) for s in secs), json.dumps([s["rel"] for s in secs if "같은이름" in s["file"]], ensure_ascii=False))
    chk("ZIP 안 파일 식별자에 안쪽 경로가 남음", any(s["file"].endswith("::a/same.txt") for s in secs) and any(s["file"].endswith("::b/same.txt") for s in secs), json.dumps([s["file"] for s in secs], ensure_ascii=False))
    # 규정 변경표: 표의 모든 행이 읽혀야 한다(칸 안 '|' 를 '\|' 로 안 쓰면 그 행이 조용히 빠졌다 — 2026-09-30)
    rc_md = os.path.join(os.path.dirname(os.path.dirname(SCR)), "_shared", "rule_changes.md")
    if os.path.exists(rc_md):
        rows = [l for l in io.open(rc_md, encoding="utf-8") if l.startswith("|") and not l.startswith("|--") and not l.startswith("| id")]
        sys.path.insert(0, SCR)
        try:
            import importlib, wiki_extract as _we
            importlib.reload(_we)
            got = _we.load_rule_changes()
        except Exception as e:
            got = []
        chk("규정 변경표: 표의 모든 행이 읽힘(칸 안 | 는 \\| 로)", len(got) == len(rows) and len(rows) >= 9, f"{len(got)}/{len(rows)}")
        cafe = next((r for r in got if r["id"] == "meeting_receipt_cafe"), None)
        hit = ["*일부 세부내역 필요 업체: 백화점, 카페, 편의점, 제과점, 슈퍼마켓", "카페 결제는 영수증 별도 첨부 필수"]
        miss = ["카페는 영수증 첨부 필수가 아님 — 보관", "카페 영수증 첨부 필수 아님", "L3 카페를 이용할 때 리유저블컵을 지참한다", "북카페 담당자에게 요청", "식사+카페 합계가 1인당 5만원이 넘지 않아야 함"]
        chk("구값 표시: 카페 영수증 옛 규칙(세부내역 필요 업체에 카페)은 잡고, '첨부 필수 아님'·북카페·L3 카페는 안 잡음",
            bool(cafe) and all(cafe["re"].search(t) for t in hit) and not any(cafe["re"].search(t) for t in miss),
            "규칙 없음" if not cafe else json.dumps([t for t in hit if not cafe["re"].search(t)] + [t for t in miss if cafe["re"].search(t)], ensure_ascii=False))
        pre = next((r for r in got if r["id"] == "meeting_preapproval"), None)
        chk("변경표 사전결재 행: 옛 값에 '휴일'이 없고(휴일·근무지 밖 사전결재는 현행) 현행 값이 그 예외를 적음",
            bool(pre) and "휴일" not in pre["old"] and "휴일" in pre["new"] and "기본 폐지" in pre["new"], "규칙 없음" if not pre else pre["old"][:80] + " // " + pre["new"][:80])
        sp = {s["file"]: s.get("stale_values") or [] for s in secs}
        chk("구값 표시: 피험자 사례비 옛 값(12만 5천원 초과 소득세)에는 붙고, 자문료 과세최저한 12만 5천원(현행)에는 안 붙음",
            any("피험자" in x for x in sp.get("피험자.txt", [])) and not sp.get("자문료.txt"), json.dumps(sp, ensure_ascii=False)[:400])
        rc4, out4 = run([os.path.join(SCR, "wiki_extract.py"), "restale"], root)
        chk("restale: 다시 추출하지 않고 구값 표시만 다시(바뀐 것 0)", "절" in out4 and "0개 바뀜" in out4, out4[-300:])
    if have_docx:
        dt = text_of(lambda s: s["file"] == "순서.docx")
        a, b, c = dt.find("첫 문단 가나다"), dt.find("표칸 라마바"), dt.find("끝 문단 사아자")
        chk("DOCX 문단·표·문단이 문서 순서대로", 0 <= a < b < c, dt[-300:])
    # ③ 긴 표 — 행 가운데서 자르지 않고, 머리 행 + 검색어 든 행 + 줄인 행 수
    os.makedirs(os.path.join(root, "notices"))
    rows = "\n".join(f"| {i} | 일반 항목 {i} 설명 설명 설명 | {i * 1000}원 |" for i in range(1, 301))
    rows = rows.replace("| 150 | 일반 항목 150", "| 150 | 숙박 한도 항목 150")
    io.open(os.path.join(root, "notices", "260930_긴표.md"), "w", encoding="utf-8").write(
        "---\ntitle: \"긴 표 공지\"\nkind: notice\nsource_url: https://example.invalid/n\ndated: 2026-09-30\n---\n# 긴 표 공지\n\n## 숙박 기준표\n\n| 번호 | 항목 | 금액 |\n| --- | --- | --- |\n" + rows + "\n")
    rc2, out2 = run([os.path.join(SCR, "wiki_ask.py"), "숙박", "--chars", "1500", "--no-staff"], root)
    ex = out2.split("## 발췌", 1)[-1].split("## 최신 확인", 1)[0]
    tl = [l for l in ex.splitlines() if l.lstrip().startswith("|")]
    chk("wiki_ask 긴 표: 표 행을 가운데서 자르지 않음", tl and all(l.rstrip().endswith("|") for l in tl), ex[-500:])
    chk("wiki_ask 긴 표: 머리 행 + 검색어 든 행 + '(표 N행 중 M행만)'", "| 번호 | 항목 | 금액 |" in ex and "숙박 한도 항목 150" in ex and "(표 300행 중 검색어 든 1행만" in ex, ex[:800])
    chk("wiki_ask 안내에 '표 통째' 가 없음", "표 통째" not in out2 and "표는 통째" not in io.open(os.path.join(SCR, "wiki_ask.py"), encoding="utf-8").read(), "")
    # 토큰 없음 → 브라우저 checkFresh 호출(로컬 version) 안내
    rc3, out3 = run([os.path.join(SCR, "wiki_ask.py"), "첨부", "--no-staff"], root)
    chk("wiki_ask 토큰 없음: 붙여 넣을 브라우저 checkFresh 호출(로컬 version)", "window.kkWiki.checkFresh({'1001': 3})" in out3 and "fmtFresh()" in out3, out3.split("## 최신 확인", 1)[-1][:400])
finally:
    shutil.rmtree(T, ignore_errors=True)
print(f"[요약] {n - fail}/{n} PASS" + (f" — {fail} FAIL" if fail else ""))
sys.exit(1 if fail else 0)
